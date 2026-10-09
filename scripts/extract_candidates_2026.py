#!/usr/bin/env python3
"""Extract candidate list order from the reviewed Praha 6 workbook."""
from __future__ import annotations

import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/Users/admin/Model_vysledku_voleb_Praha6_opraveno.xlsx")
TARGET = ROOT / "data" / "candidates-2026.json"

workbook = load_workbook(SOURCE, read_only=True, data_only=False)
sheet = workbook["Kandidáti"]
candidates = []
for list_number, list_name, order, name, *_ in sheet.iter_rows(min_row=6, values_only=True):
    if not isinstance(list_number, int) or not isinstance(order, int) or not isinstance(name, str):
        continue
    candidates.append({"list_id": f"l{list_number}", "list_name": list_name, "ballot_order": order, "name": name})

if len(candidates) != 472:
    raise ValueError(f"Expected 472 candidates, got {len(candidates)}")
TARGET.write_text(json.dumps({"source": SOURCE.name, "candidates": candidates}, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
print(f"{len(candidates)} candidates -> {TARGET}")
